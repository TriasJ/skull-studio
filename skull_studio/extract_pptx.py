"""PPTX import front-end for the pipeline.

Renders slide backgrounds via LibreOffice (pptx -> pdf -> PNG with PyMuPDF),
runs MinerU OCR on the rendered slides to recover text baked into pictures, and
reads shape geometry/text with python-pptx. The result:
  - large pictures covering most of a slide  -> background (full-slide render)
  - text (OCR'd from images, or real PPTX text frames) -> text elements
  - other pictures / shapes                  -> image elements
Text elements carry real text, so the PPTX exporter can emit editable text boxes.

Usage:
  python -m skull_studio.extract_pptx render <deck.pptx>   # LibreOffice + PyMuPDF
  python -m skull_studio.extract_pptx mineru <deck.pptx>   # OCR the rendered slides
  python -m skull_studio.extract_pptx draft  <deck.pptx>   # merge OCR + shapes -> manifest
"""
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract  # reuse render(), run_mineru(), draft()
import extract_model3d  # raw-XML reader for inserted 3D models (am3d:model3d)

ROOT = Path(__file__).resolve().parent.parent
WORK = ROOT / "work"
BG_AREA = 0.60          # picture area fraction that counts as the background
FULL_SHAPE_AREA = 0.92  # a non-picture shape this big is a full-bleed backdrop


def _pdf_for(pptx_path: Path) -> Path:
    return WORK / "pptx" / (pptx_path.stem + ".pdf")


def find_soffice():
    """Locate LibreOffice across platforms (PATH first, then common installs)."""
    candidates = [shutil.which("soffice"), shutil.which("libreoffice"),
                  # Windows
                  r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                  # macOS
                  "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                  # Linux
                  "/usr/bin/soffice", "/usr/bin/libreoffice",
                  "/snap/bin/libreoffice", "/opt/libreoffice/program/soffice"]
    for p in candidates:
        if p and Path(p).exists():
            return p
    return None


def render(pptx_path: Path) -> None:
    soffice = find_soffice()
    if not soffice:
        sys.exit("LibreOffice (soffice) not found - needed to render PPTX slides.")
    out = WORK / "pptx"
    out.mkdir(parents=True, exist_ok=True)
    print("Converting PPTX -> PDF via LibreOffice...")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir",
                    str(out), str(pptx_path)], check=True, timeout=300)
    pdf = _pdf_for(pptx_path)
    if not pdf.exists():
        sys.exit(f"LibreOffice did not produce {pdf}")
    extract.render(pdf)


def mineru(pptx_path: Path) -> None:
    pdf = _pdf_for(pptx_path)
    if not pdf.exists():
        sys.exit("Run the render stage first (need the converted PDF).")
    extract.run_mineru(pdf)


_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _prune_unsupported(slide):
    """Remove spTree children python-pptx's shape factory cannot build.

    Inserted 3D models live in ``mc:AlternateContent`` and ink lives in
    ``p:contentPart``; iterating either raises ``has_ph_elm`` and aborts the
    whole slide. We strip them from the in-memory tree only (the file is never
    written) — 3D models are read from the raw package separately, and ink is
    not imported. Returns the count removed."""
    spTree = slide.shapes._spTree
    removed = 0
    for child in list(spTree):
        if child.tag in (f"{{{_MC}}}AlternateContent", f"{{{_P}}}contentPart"):
            spTree.remove(child)
            removed += 1
    return removed


def _norm(shape, SW, SH):
    x0, y0 = max(0, shape.left / SW), max(0, shape.top / SH)
    x1 = min(1, (shape.left + shape.width) / SW)
    y1 = min(1, (shape.top + shape.height) / SH)
    return [round(x0, 4), round(y0, 4), round(x1, 4), round(y1, 4)]


def _overlaps(a, b):
    ix = max(0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    if inter <= 0:
        return 0.0
    ua = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    return inter / ua if ua else 0.0


def draft(pptx_path: Path) -> None:
    pdf = _pdf_for(pptx_path)
    # OCR-based draft when MinerU ran; otherwise keep the existing (minimal)
    # manifest and rely on PPTX shapes alone
    has_ocr = (WORK / "mineru").exists() and any((WORK / "mineru").rglob("*content_list.json"))
    if has_ocr:
        extract.draft(pdf)  # text + image regions detected on the rendered slides
    m = json.loads((WORK / "manifest.json").read_text(encoding="utf-8"))

    # rebrand to the PPTX source
    m["meta"] = {"title": re.sub(r"[_-]+", " ", pptx_path.stem),
                 "sourcePptx": pptx_path.name,
                 "docId": "deck-" + hashlib.sha1(pptx_path.read_bytes()).hexdigest()[:8]}

    prs = Presentation(str(pptx_path))
    SW, SH = prs.slide_width, prs.slide_height
    slides = list(prs.slides)
    added_pics = added_txt = exact_txt = removed_dup = 0

    for i, msl in enumerate(m["slides"]):
        for el in msl["elements"]:                      # ensure new fields exist
            el.setdefault("blendMode", None)
            el.setdefault("opacity", 1)
            el.setdefault("hidden", False)
        if i >= len(slides):
            continue
        _prune_unsupported(slides[i])   # strip 3D-model/ink elems python-pptx can't build
        shapes = [s for s in slides[i].shapes if s.left is not None and s.width]
        zmax = max([e.get("z", 0) for e in msl["elements"]] + [0])
        pptx_imgs = []

        for shape in shapes:
            bbox = _norm(shape, SW, SH)
            if bbox[2] - bbox[0] < 0.01 or bbox[3] - bbox[1] < 0.01:
                continue
            area = (bbox[2] - bbox[0]) * (bbox[3] - bbox[1])
            is_pic = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            text = shape.text.strip() if shape.has_text_frame else ""

            if text:
                # real PPTX text box: replace overlapping OCR text with exact text
                hit = max((e for e in msl["elements"] if e["type"] == "text"),
                          key=lambda e: _overlaps(bbox, e["bbox"]), default=None)
                if hit and _overlaps(bbox, hit["bbox"]) > 0.25:
                    hit["text"] = text
                    exact_txt += 1
                else:
                    zmax += 1
                    msl["elements"].append(_text_el(msl["id"], zmax, bbox, text))
                    added_txt += 1
                continue
            # keep EVERY non-text shape (pictures, freeforms, autoshapes) except
            # the full-slide background - cropped from the render so vectors and
            # effects are preserved as they appear on the slide
            if is_pic and area > BG_AREA:
                continue                                 # background picture
            if not is_pic and area > FULL_SHAPE_AREA:
                continue                                 # full-bleed backdrop shape
            zmax += 1
            el = _img_el(msl["id"], zmax, bbox)
            msl["elements"].append(el)
            pptx_imgs.append(el)
            added_pics += 1

        # drop MinerU image regions that an exact PPTX picture already covers
        for oi in [e for e in msl["elements"] if e["type"] == "image" and e not in pptx_imgs]:
            if any(_overlaps(oi["bbox"], pe["bbox"]) > 0.5 for pe in pptx_imgs):
                msl["elements"].remove(oi)
                removed_dup += 1

    # inserted 3D models (invisible to python-pptx; read from raw slide XML)
    added_models = _add_models(m, pptx_path)

    (WORK / "manifest.json").write_text(
        json.dumps(m, indent=2, ensure_ascii=False), encoding="utf-8")
    n = sum(len(s["elements"]) for s in m["slides"])
    print(f"Merged PPTX shapes: +{exact_txt} exact text, +{added_txt} text boxes, "
          f"+{added_pics} shapes/pictures (-{removed_dup} OCR dups), "
          f"+{added_models} 3D models. "
          f"Total {n} elements across {len(m['slides'])} slides.")


def _add_models(m, pptx_path):
    """Detect inserted 3D models and append them as `model3d` elements.

    The visible crop still comes from the LibreOffice render (which rasterizes
    each model's 2D fallback preview), so a model already works as a static
    element with no runtime support; the extra fields enable real 3D later and
    a loss-free round-trip back to PPTX.
    """
    try:
        models_by_slide = extract_model3d.parse(pptx_path)
    except Exception as e:                                   # never break import
        print(f"  (3D model scan skipped: {e})")
        return 0
    if not models_by_slide:
        return 0
    models_dir = WORK / "models"
    models_dir.mkdir(parents=True, exist_ok=True)
    written, n = set(), 0
    with zipfile.ZipFile(str(pptx_path)) as z:
        names = set(z.namelist())
        for i, msl in enumerate(m["slides"]):
            for info in models_by_slide.get(i, []):
                if not info.get("bbox"):
                    continue
                z_next = max([e.get("z", 0) for e in msl["elements"]] + [0]) + 1
                msl["elements"].append(_model3d_el(msl["id"], z_next, info))
                for part in (info.get("glb_part"), info.get("preview_part")):
                    if part and part in names and part not in written:
                        (models_dir / part.rsplit("/", 1)[-1]).write_bytes(z.read(part))
                        written.add(part)
                n += 1
            # OCR image regions hidden behind a model are duplicates of its preview
            mboxes = [info["bbox"] for info in models_by_slide.get(i, []) if info.get("bbox")]
            for oi in [e for e in msl["elements"] if e["type"] == "image"]:
                if any(_overlaps(oi["bbox"], mb) > 0.5 for mb in mboxes):
                    msl["elements"].remove(oi)
    return n


def _base(sid, z, bbox, etype, role, text=""):
    eid = f"{sid}_p{z:02d}"
    return {"id": eid, "type": etype, "role": role, "text": text[:300],
            "bbox": bbox, "crop": f"crops/{eid}.webp", "cropBbox": None,
            "z": z, "cleanup": "none", "fillColor": None, "patch": None,
            "entrance": {"type": "fadeIn", "delay": 0.2, "duration": 0.8, "ease": "power2.out"},
            "idle": [], "parallax": 0.0, "lines": None, "rig": None,
            "blendMode": None, "opacity": 1, "hidden": False}


def _text_el(sid, z, bbox, text):
    return _base(sid, z, bbox, "text", "title" if bbox[1] < 0.18 else "body", text)


def _img_el(sid, z, bbox):
    return _base(sid, z, bbox, "image", "figure")


def _model3d_el(sid, z, info):
    """Build a `model3d` element from an extract_model3d.parse() entry.

    `crop` (the rendered preview) is the Tier-1 display + universal fallback;
    `modelSrc`/`model3d` enable three.js rendering and PPTX round-trip.
    """
    el = _base(sid, z, info["bbox"], "model3d", "figure", info.get("descr") or "")
    el["name"] = info.get("name")
    el["parallax"] = 0.04                                    # models read as foreground
    base = lambda part: part.rsplit("/", 1)[-1] if part else None
    el["modelSrc"] = f"models/{base(info['glb_part'])}" if info.get("glb_part") else None
    anim = info.get("anim") or {}
    el["model3d"] = {
        "clip": anim.get("clip", 0) if anim else None,
        "loop": anim.get("loop", False),
        "durationMs": anim.get("durationMs"),
        "autoRotate": 0,                                     # editor-added emphasis (deg/s)
        "camera": info.get("camera"),
        "transform": info.get("transform"),
        "previewSrc": f"models/{base(info['preview_part'])}" if info.get("preview_part") else None,
        "sourceXml": info.get("source_xml"),                # verbatim, for round-trip export
    }
    return el


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] not in ("render", "mineru", "draft"):
        sys.exit(__doc__)
    stage, pptx = sys.argv[1], Path(sys.argv[2]).resolve()
    {"render": render, "mineru": mineru, "draft": draft}[stage](pptx)
